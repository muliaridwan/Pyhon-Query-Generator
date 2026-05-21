from __future__ import annotations

import html
import json
import os
import tempfile
import uuid
from collections import OrderedDict
from datetime import datetime, date
from decimal import Decimal
from typing import Any

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, PlainTextResponse
from openpyxl import load_workbook
import uvicorn

HOST = "0.0.0.0"
PORT = 5000
app = FastAPI(title="Excel Delete Insert SQL Generator")
PARSED_UPLOADS: dict[str, dict[str, Any]] = {}


def normalize_value(value: Any) -> Any:
    if isinstance(value, str):
        text = value.strip()
        if text == "" or text.upper() == "NULL":
            return None
        return text
    return value


def excel_to_python_value(value: Any) -> Any:
    value = normalize_value(value)
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    return value


def sql_literal(value: Any) -> str:
    value = excel_to_python_value(value)
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, datetime):
        return "'" + value.strftime("%Y-%m-%d %H:%M:%S.%f")[:-3] + "'"
    return "'" + str(value).replace("'", "''") + "'"


def build_match_expression(alias: str, fields: list[str], row: dict[str, Any], joiner: str = "AND", skip_nulls: bool = False) -> str:
    parts = []
    for field in fields:
        value = excel_to_python_value(row.get(field))
        if value is None:
            if skip_nulls:
                continue
            parts.append(f"{alias}.[{field}] IS NULL")
        else:
            parts.append(f"{alias}.[{field}] = {sql_literal(value)}")
    return f" {joiner} ".join(parts) if parts else ""


def distinct_key_rows(rows: list[dict[str, Any]], key_fields: list[str]) -> list[dict[str, Any]]:
    seen = set()
    output = []
    for row in rows:
        key = tuple(excel_to_python_value(row.get(field)) for field in key_fields)
        if key in seen:
            continue
        seen.add(key)
        output.append(row)
    return output


def is_effectively_empty_row(row: dict[str, Any], fields: list[str]) -> bool:
    return all(excel_to_python_value(row.get(field)) is None for field in fields)


def build_delete_sql(table_name: str, rows: list[dict[str, Any]], delete_fields: list[str]) -> str:
    if not delete_fields:
        return f"-- DELETE skipped for [{table_name}] because no delete key selected."
    keys = distinct_key_rows(rows, delete_fields)
    predicates = [build_match_expression("T", delete_fields, row, joiner="OR") for row in keys]
    return (
        f"-- DELETE [{table_name}]\n"
        f"DELETE T\n"
        f"FROM [{table_name}] T\n"
        f"WHERE " + "\n    OR ".join(predicates) + ";"
    )


def build_insert_sql(table_name: str, rows: list[dict[str, Any]], unique_fields: list[str], all_fields: list[str]) -> str:
    if not rows:
        return f"-- INSERT skipped for [{table_name}] because there is no row data."
    rows_for_insert = distinct_key_rows(rows, unique_fields) if unique_fields else rows
    columns = [field for field in all_fields if any(excel_to_python_value(row.get(field)) is not None for row in rows)]
    if not columns:
        return f"-- INSERT skipped for [{table_name}] because all values are NULL."
    col_sql = ", ".join(f"[{c}]" for c in columns)
    chunks = [f"-- INSERT [{table_name}]\n"]
    for idx, row in enumerate(rows_for_insert, start=1):
        values_sql = ", ".join(sql_literal(row.get(c)) for c in columns)
        if unique_fields:
            match_expr = build_match_expression("T", unique_fields, row)
            stmt = (
                f"IF NOT EXISTS (SELECT 1 FROM [{table_name}] T WHERE {match_expr})\n"
                f"BEGIN\n"
                f"    INSERT INTO [{table_name}] ({col_sql})\n"
                f"    VALUES ({values_sql});\n"
                f"END;"
            )
        else:
            stmt = f"INSERT INTO [{table_name}] ({col_sql})\nVALUES ({values_sql});"
        chunks.append(stmt)
        if idx != len(rows_for_insert):
            chunks.append("")
    return "\n".join(chunks)


def parse_compound_sheet(filename: str) -> dict[str, Any]:
    wb = load_workbook(filename=filename, data_only=False)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel kosong")
    header = [str(v).strip() if v is not None else "" for v in rows[0]]
    starts = [i for i, v in enumerate(header) if v == "TableName"]
    if not starts:
        raise ValueError("Header TableName tidak ditemukan")
    starts.append(len(header))
    segments = []
    for i in range(len(starts) - 1):
        s, e = starts[i], starts[i + 1]
        seg_headers = header[s:e]
        if len(seg_headers) < 2:
            continue
        segments.append({"start": s, "end": e, "fields": seg_headers[1:]})

    tables: dict[str, dict[str, Any]] = OrderedDict()
    for row in rows[1:]:
        for seg in segments:
            values = list(row[seg["start"]:seg["end"]])
            if not values:
                continue
            table_name = normalize_value(values[0])
            if not table_name:
                continue
            fields = seg["fields"]
            payload = values[1:]
            row_dict = {
                field: excel_to_python_value(payload[idx]) if idx < len(payload) else None
                for idx, field in enumerate(fields)
            }
            if table_name not in tables:
                tables[table_name] = {"name": table_name, "fields": list(fields), "rows": []}
            else:
                for field in fields:
                    if field not in tables[table_name]["fields"]:
                        tables[table_name]["fields"].append(field)
            tables[table_name]["rows"].append(row_dict)
    return {"tables": list(tables.values()), "source_rows": max(len(rows) - 1, 0)}


def serialize_parsed(parsed: dict[str, Any]) -> dict[str, Any]:
    return json.loads(json.dumps(parsed, default=str))


def h(text: Any) -> str:
    return html.escape("" if text is None else str(text))


def render_page(error: str | None = None, upload: dict[str, Any] | None = None, result: dict[str, Any] | None = None, selected_delete: dict[str, list[str]] | None = None, selected_unique: dict[str, list[str]] | None = None) -> str:
    parts = ["""<!doctype html>
<html lang='en'>
<head>
<meta charset='utf-8'>
<meta name='viewport' content='width=device-width, initial-scale=1'>
<title>Excel Delete Insert SQL Generator</title>
<style>
body { font-family: Arial, sans-serif; margin: 24px; background: #f6f7fb; color: #1d2433; }
.container { max-width: 1400px; margin: 0 auto; }
.card { background: white; border-radius: 12px; padding: 18px; box-shadow: 0 2px 10px rgba(0,0,0,0.08); margin-bottom: 18px; }
.grid { display: grid; grid-template-columns: 1fr; gap: 16px; }
.table-card { border: 1px solid #e1e5ee; border-radius: 12px; padding: 14px; background: #fcfcfe; }
.field-list { max-height: 260px; overflow: auto; border: 1px solid #e6e9f2; border-radius: 8px; padding: 10px; background: white; }
label { display: block; margin: 4px 0; }
.field-grid { display: grid; grid-template-columns: 1fr; gap: 14px; }
.field-section { min-width: 0; }
button { border: none; background: #2d63ff; color: white; padding: 10px 16px; border-radius: 8px; cursor: pointer; font-weight: 600; }
button.secondary { background: #5b6478; }
textarea { width: 100%; min-height: 520px; font-family: Consolas, monospace; font-size: 13px; border: 1px solid #d7dcea; border-radius: 10px; padding: 12px; box-sizing: border-box; }
.pill { display: inline-block; padding: 4px 10px; background: #edf2ff; color: #2346a3; border-radius: 999px; font-size: 12px; margin-right: 6px; margin-bottom: 6px; }
.small { font-size: 12px; color: #586074; }
.warn { background: #fff8e6; border: 1px solid #f4d37a; padding: 10px 12px; border-radius: 8px; }
.actions { display:flex; gap:10px; flex-wrap:wrap; align-items:center; }
</style>
</head>
<body><div class='container'>
<div class='card'>
<h1>Excel Delete Insert SQL Generator</h1>
<p class='small'>Upload file Excel, pilih field delete key dan unique / primary key insert, lalu aplikasi membuat SQL Server script delete + insert. Insert tidak duplicate berdasarkan key unik yang dipilih.</p>
<form method='post' action='/upload' enctype='multipart/form-data'>
  <input type='file' name='excel_file' accept='.xlsx,.xlsm' required>
  <button type='submit'>Upload dan Baca Excel</button>
</form>
</div>
"""]
    if error:
        parts.append(f"<div class='card warn'>{h(error)}</div>")
    if upload:
        parts.append(f"""
<form method='post' action='/process/{h(upload['id'])}'>
<div class='card'>
  <h2>File berhasil dibaca</h2>
  <p class='small'>Upload ID: <code>{h(upload['id'])}</code></p>
  <p><span class='pill'>Total tabel: {h(upload['table_count'])}</span><span class='pill'>Total row sumber: {h(upload['source_rows'])}</span></p>
  <p class='small'>Centang field pada kolom kiri untuk delete key, dan kolom kanan untuk unique / primary key insert.</p>
</div>
<div class='grid'>
""")
        for table in upload["tables"]:
            parts.append(f"""
<div class='table-card'>
  <h3>{h(table['name'])}</h3>
  <p class='small'>Field: {h(table['field_count'])} | Row data: {h(table['row_count'])}</p>
  <div class='field-grid'>
    <div class='field-section'>
      <strong>Delete key</strong>
      <div class='field-list'>
""")
            selected_delete_set = set((selected_delete or {}).get(table['name'], []))
            selected_unique_set = set((selected_unique or {}).get(table['name'], []))
            for field in table["fields"]:
                checked = " checked" if field in selected_delete_set else ""
                parts.append(f"<label><input type='checkbox' name='delete__{h(table['name'])}' value='{h(field)}'{checked}> {h(field)}</label>")
            parts.append("</div></div><div class='field-section'><strong>Unique / Primary key insert</strong><div class='field-list'>")
            for field in table["fields"]:
                checked = " checked" if field in selected_unique_set else ""
                parts.append(f"<label><input type='checkbox' name='unique__{h(table['name'])}' value='{h(field)}'{checked}> {h(field)}</label>")
            parts.append("</div></div></div></div>")
        parts.append("</div><div class='card actions'><button type='submit'>Process SQL</button><a href='/'><button type='button' class='secondary'>Reset</button></a></div></form>")
    if result:
        parts.append(f"""
<div class='card'>
  <h2>SQL berhasil dibuat</h2>
  <div class='actions'>
    <a href='/download/{h(result['upload_id'])}'><button type='button'>Download .sql</button></a>
    <a href='/'><button type='button' class='secondary'>Upload file lain</button></a>
  </div>
  <p class='small'>Preview maksimum 200.000 karakter pertama.</p>
  <textarea readonly>{h(result['preview'])}</textarea>
</div>
""")
    parts.append("</div></body></html>")
    return "".join(parts)


@app.get("/", response_class=HTMLResponse)
def index() -> str:
    return render_page()


@app.post("/upload", response_class=HTMLResponse)
async def upload(excel_file: UploadFile = File(...)) -> str:
    if not excel_file.filename:
        return render_page(error="File Excel wajib dipilih.")
    suffix = os.path.splitext(excel_file.filename)[1] or ".xlsx"
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(await excel_file.read())
        tmp_path = tmp.name
    try:
        parsed_excel = parse_compound_sheet(tmp_path)
    except Exception as exc:
        return render_page(error=f"Gagal membaca Excel: {exc}")
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass

    upload_id = uuid.uuid4().hex
    PARSED_UPLOADS[upload_id] = serialize_parsed(parsed_excel)
    upload_meta = {
        "id": upload_id,
        "table_count": len(parsed_excel["tables"]),
        "source_rows": parsed_excel["source_rows"],
        "tables": [
            {"name": t["name"], "fields": t["fields"], "field_count": len(t["fields"]), "row_count": len(t["rows"])}
            for t in parsed_excel["tables"]
        ],
    }
    return render_page(upload=upload_meta)


@app.post("/process/{upload_id}", response_class=HTMLResponse)
async def process(upload_id: str, request: Request) -> str:
    stored = PARSED_UPLOADS.get(upload_id)
    if not stored:
        return render_page(error="Session upload tidak ditemukan. Upload ulang file Excel.")

    form = await request.form()
    sql_parts = [
        "/* Generated by Excel Delete Insert SQL Generator */",
        "SET NOCOUNT ON;",
        "BEGIN TRY",
        "    BEGIN TRAN;",
        "",
    ]
    delete_parts: list[str] = ["-- ===== DELETE SECTION =====", ""]
    insert_parts: list[str] = ["-- ===== INSERT SECTION =====", ""]
    upload_meta = {
        "id": upload_id,
        "table_count": len(stored["tables"]),
        "source_rows": stored["source_rows"],
        "tables": [],
    }
    selected_delete: dict[str, list[str]] = {}
    selected_unique: dict[str, list[str]] = {}
    for table in stored["tables"]:
        name = table["name"]
        fields = table["fields"]
        rows = table["rows"]
        delete_fields = form.getlist(f"delete__{name}")
        unique_fields = form.getlist(f"unique__{name}")
        selected_delete[name] = delete_fields
        selected_unique[name] = unique_fields

        upload_meta["tables"].append({"name": name, "fields": fields, "field_count": len(fields), "row_count": len(rows)})

        header_parts = [
            "-- ============================================",
            f"-- TABLE [{name}]",
            f"-- Delete keys : {', '.join(delete_fields) if delete_fields else '(none)'}",
            f"-- Insert unique keys : {', '.join(unique_fields) if unique_fields else '(none)'}",
            "-- ============================================",
        ]

        delete_parts.extend(header_parts)
        delete_parts.append(build_delete_sql(name, rows, delete_fields))
        delete_parts.append("")

        insert_parts.extend(header_parts)
        insert_parts.append(build_insert_sql(name, rows, unique_fields, fields))
        insert_parts.append("")

    sql_parts.extend(delete_parts)
    sql_parts.append("")
    sql_parts.extend(insert_parts)
    sql_parts.extend([
        "    COMMIT TRAN;",
        "END TRY",
        "BEGIN CATCH",
        "    IF @@TRANCOUNT > 0 ROLLBACK TRAN;",
        "    THROW;",
        "END CATCH;",
    ])
    final_sql = "\n".join(sql_parts)
    PARSED_UPLOADS[upload_id]["generated_sql"] = final_sql
    return render_page(upload=upload_meta, result={"upload_id": upload_id, "preview": final_sql[:200000]}, selected_delete=selected_delete, selected_unique=selected_unique)


@app.get("/download/{upload_id}", response_class=PlainTextResponse)
def download_sql(upload_id: str) -> PlainTextResponse:
    stored = PARSED_UPLOADS.get(upload_id)
    if not stored or "generated_sql" not in stored:
        raise HTTPException(status_code=404, detail="SQL belum tersedia")
    return PlainTextResponse(
        stored["generated_sql"],
        media_type="application/sql",
        headers={"Content-Disposition": f"attachment; filename=generated_{upload_id}.sql"},
    )


if __name__ == "__main__":
    uvicorn.run(app, host=HOST, port=PORT)
